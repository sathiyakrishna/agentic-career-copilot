import json
from io import BytesIO

from dotenv import load_dotenv
from openai import OpenAI
from docx import Document


load_dotenv()

client = OpenAI()


# =========================================================
# RESUME AGENT
# =========================================================

def run_resume_agent(
    job_description,
    resume_text,
    evidence,
    fit_result
):
    """
    Tailor the candidate's master resume to the JD
    using only verified candidate evidence.
    """

    evidence_text = "\n\n".join(
        [
            f"EVIDENCE {i + 1}:\n{chunk}"
            for i, chunk in enumerate(evidence)
        ]
    )

    system_prompt = """
You are the Resume Agent for an AI Career Copilot.

Your task is to tailor the candidate's existing resume
to the supplied job description.

STRICT EVIDENCE RULES:

1. Use ONLY facts contained in the master resume
   and verified candidate evidence.

2. NEVER invent:
   - employers
   - roles
   - employment dates
   - responsibilities
   - projects
   - achievements
   - metrics
   - certifications
   - technologies
   - skills

3. You may rephrase existing experience.

4. You may reorder existing experience.

5. You may emphasize relevant existing achievements.

6. Preserve the original factual meaning.

7. Never add a JD keyword merely for ATS optimization
   unless the candidate evidence genuinely supports it.

8. Unsupported JD requirements must be excluded from
   the resume and listed under unsupported_requirements.

9. Do not exaggerate seniority.

10. Keep the resume ATS friendly.

Return ONLY valid JSON.

Use exactly this structure:

{
  "professional_summary": "",
  "skills": [],
  "experience": [
    {
      "company": "",
      "role": "",
      "dates": "",
      "bullets": []
    }
  ],
  "projects": [
    {
      "name": "",
      "bullets": []
    }
  ],
  "education": [],
  "certifications": [],
  "unsupported_requirements": []
}
"""

    user_prompt = f"""
JOB DESCRIPTION:

{job_description}


MASTER RESUME:

{resume_text}


VERIFIED RAG EVIDENCE:

{evidence_text}


FIT ANALYSIS:

{json.dumps(fit_result, indent=2)}


Create a tailored resume for this job.

The output must remain completely faithful
to the candidate's master resume.
"""

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        temperature=0,
        response_format={
            "type": "json_object"
        },
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ]
    )

    result_text = (
        response
        .choices[0]
        .message
        .content
    )

    return json.loads(result_text)


# =========================================================
# DOCX GENERATOR
# =========================================================

def create_resume_docx(resume_data):
    """
    Convert tailored resume JSON into
    an ATS-friendly Word document.
    """

    document = Document()

    # Professional Summary
    document.add_heading(
        "Professional Summary",
        level=1
    )

    document.add_paragraph(
        resume_data.get(
            "professional_summary",
            ""
        )
    )

    # Skills
    document.add_heading(
        "Core Skills",
        level=1
    )

    skills = resume_data.get(
        "skills",
        []
    )

    document.add_paragraph(
        " | ".join(skills)
    )

    # Experience
    document.add_heading(
        "Professional Experience",
        level=1
    )

    for item in resume_data.get(
        "experience",
        []
    ):
        role = item.get(
            "role",
            ""
        )

        company = item.get(
            "company",
            ""
        )

        document.add_heading(
            f"{role} — {company}",
            level=2
        )

        dates = item.get(
            "dates",
            ""
        )

        if dates:
            document.add_paragraph(
                dates
            )

        for bullet in item.get(
            "bullets",
            []
        ):
            document.add_paragraph(
                bullet,
                style="List Bullet"
            )

    # Projects
    projects = resume_data.get(
        "projects",
        []
    )

    if projects:
        document.add_heading(
            "Projects",
            level=1
        )

        for project in projects:
            document.add_heading(
                project.get(
                    "name",
                    ""
                ),
                level=2
            )

            for bullet in project.get(
                "bullets",
                []
            ):
                document.add_paragraph(
                    bullet,
                    style="List Bullet"
                )

    # Education
    education = resume_data.get(
        "education",
        []
    )

    if education:
        document.add_heading(
            "Education",
            level=1
        )

        for item in education:
            document.add_paragraph(
                str(item)
            )

    # Certifications
    certifications = resume_data.get(
        "certifications",
        []
    )

    if certifications:
        document.add_heading(
            "Certifications",
            level=1
        )

        for item in certifications:
            document.add_paragraph(
                str(item),
                style="List Bullet"
            )

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output