import hashlib
import json
from pathlib import Path
from io import BytesIO

import chromadb
import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI
from pypdf import PdfReader
from docx import Document


# =========================================================
# CONFIGURATION
# =========================================================

load_dotenv()

client = OpenAI()

RESUME_DIR = Path("data/resume")
RESUME_DIR.mkdir(parents=True, exist_ok=True)

MASTER_RESUME_PATH = RESUME_DIR / "master_resume.pdf"
CHROMA_PATH = "data/chroma"


# =========================================================
# CHROMADB
# =========================================================

chroma_client = chromadb.PersistentClient(
    path=CHROMA_PATH
)

resume_collection = chroma_client.get_or_create_collection(
    name="candidate_resume"
)


# =========================================================
# RESUME FUNCTIONS
# =========================================================

def extract_pdf_text(pdf_file):
    """
    Extract text from a PDF resume.
    """

    reader = PdfReader(pdf_file)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text.strip()


def save_master_resume(uploaded_file):
    """
    Save uploaded resume as active master resume.
    """

    with open(MASTER_RESUME_PATH, "wb") as f:
        f.write(uploaded_file.getbuffer())


def resume_exists():
    """
    Check whether master resume exists.
    """

    return MASTER_RESUME_PATH.exists()


def get_resume_hash():
    """
    Generate version ID for active resume.
    """

    if not resume_exists():
        return None

    with open(MASTER_RESUME_PATH, "rb") as f:
        file_bytes = f.read()

    return hashlib.sha256(
        file_bytes
    ).hexdigest()[:12]


# =========================================================
# RAG — CHUNKING
# =========================================================

def chunk_text(
    text,
    chunk_size=180,
    overlap=30
):
    """
    Split resume into overlapping chunks.
    """

    words = text.split()

    chunks = []

    start = 0

    while start < len(words):

        end = start + chunk_size

        chunk = " ".join(
            words[start:end]
        )

        if chunk.strip():
            chunks.append(chunk)

        start += chunk_size - overlap

    return chunks


# =========================================================
# RAG — INDEXING
# =========================================================

def index_resume(
    resume_text,
    resume_version
):
    """
    Embed resume chunks and persist in ChromaDB.
    """

    chunks = chunk_text(resume_text)

    existing = resume_collection.get()

    existing_metadata = existing.get(
        "metadatas",
        []
    )

    existing_versions = set()

    for metadata in existing_metadata:

        if metadata:

            version = metadata.get(
                "resume_version"
            )

            if version:
                existing_versions.add(
                    version
                )

    # Resume already indexed
    if resume_version in existing_versions:

        return len(chunks)

    # Remove old resume embeddings
    if existing.get("ids"):

        resume_collection.delete(
            ids=existing["ids"]
        )

    # Generate embeddings
    for i, chunk in enumerate(chunks):

        response = client.embeddings.create(
            model="text-embedding-3-small",
            input=chunk
        )

        embedding = (
            response.data[0].embedding
        )

        resume_collection.add(

            ids=[
                f"{resume_version}_{i}"
            ],

            documents=[
                chunk
            ],

            embeddings=[
                embedding
            ],

            metadatas=[
                {
                    "resume_version":
                        resume_version,

                    "chunk_number":
                        i
                }
            ]
        )

    return len(chunks)


# =========================================================
# RAG — RETRIEVAL
# =========================================================

def retrieve_resume_evidence(
    job_description,
    top_k=6
):
    """
    Retrieve resume evidence relevant to the JD.
    """

    response = client.embeddings.create(
        model="text-embedding-3-small",
        input=job_description
    )

    query_embedding = (
        response.data[0].embedding
    )

    total_chunks = resume_collection.count()

    if total_chunks == 0:
        return []

    actual_top_k = min(
        top_k,
        total_chunks
    )

    results = resume_collection.query(

        query_embeddings=[
            query_embedding
        ],

        n_results=actual_top_k
    )

    documents = results.get(
        "documents",
        [[]]
    )

    if not documents:
        return []

    return documents[0]


# =========================================================
# FIT AGENT
# =========================================================

def run_fit_agent(
    job_description,
    evidence
):

    evidence_text = "\n\n".join(
        [
            f"EVIDENCE {i + 1}:\n{chunk}"
            for i, chunk in enumerate(evidence)
        ]
    )

    system_prompt = """
You are the Fit Agent for an AI Career Copilot.

Evaluate how well VERIFIED candidate resume evidence
matches a job description.

STRICT RULES:

1. Use ONLY candidate evidence supplied to you.
2. Never invent experience.
3. Never invent skills.
4. Never invent certifications.
5. Never invent achievements or metrics.
6. Never assume experience merely because two concepts
   are semantically related.
7. Unsupported requirements must be classified as GAP.

Classify requirements as:

STRONG MATCH
PARTIAL MATCH
GAP

Return ONLY valid JSON.

Use exactly this structure:

{
  "fit_score": 0,
  "summary": "",
  "strong_matches": [
    {
      "requirement": "",
      "evidence": "",
      "reason": ""
    }
  ],
  "partial_matches": [
    {
      "requirement": "",
      "evidence": "",
      "reason": ""
    }
  ],
  "gaps": [
    {
      "requirement": "",
      "reason": ""
    }
  ],
  "recommendation": ""
}

fit_score must be an integer from 0 to 100.
"""

    user_prompt = f"""
JOB DESCRIPTION:

{job_description}


VERIFIED CANDIDATE EVIDENCE:

{evidence_text}


Evaluate the candidate's fit for this job.

Do not assume unsupported experience.
"""

    response = client.chat.completions.create(

        model="gpt-4o-mini",

        temperature=0,

        response_format={
            "type": "json_object"
        },

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ]
    )

    result_text = (
        response
        .choices[0]
        .message
        .content
    )

    return json.loads(
        result_text
    )


# =========================================================
# RESUME AGENT
# =========================================================

def run_resume_agent(
    job_description,
    resume_text,
    evidence,
    fit_result
):

    evidence_text = "\n\n".join(
        [
            f"EVIDENCE {i + 1}:\n{chunk}"
            for i, chunk in enumerate(evidence)
        ]
    )

    system_prompt = """
You are the Resume Agent for an AI Career Copilot.

Your task is to tailor the candidate's existing resume
to the supplied job description.

STRICT EVIDENCE RULES:

1. Use ONLY facts contained in the master resume
   and verified candidate evidence.

2. NEVER invent:
   - employers
   - roles
   - employment dates
   - responsibilities
   - projects
   - achievements
   - metrics
   - certifications
   - technologies
   - skills

3. You may rephrase existing experience.

4. You may reorder existing experience.

5. You may emphasize relevant existing achievements.

6. Preserve the original factual meaning.

7. Never add a JD keyword merely for ATS optimization
   unless the candidate evidence genuinely supports it.

8. Unsupported JD requirements must be excluded from
   the resume and listed under unsupported_requirements.

9. Do not exaggerate seniority.

10. Keep the resume ATS friendly.

Return ONLY valid JSON.

Use exactly this structure:

{
  "professional_summary": "",
  "skills": [],
  "experience": [
    {
      "company": "",
      "role": "",
      "dates": "",
      "bullets": []
    }
  ],
  "projects": [
    {
      "name": "",
      "bullets": []
    }
  ],
  "education": [],
  "certifications": [],
  "unsupported_requirements": []
}
"""

    user_prompt = f"""
JOB DESCRIPTION:

{job_description}


MASTER RESUME:

{resume_text}


VERIFIED RAG EVIDENCE:

{evidence_text}


FIT ANALYSIS:

{json.dumps(fit_result, indent=2)}


Create a tailored resume for this job.

The output must remain completely faithful
to the candidate's master resume.
"""

    response = client.chat.completions.create(

        model="gpt-4o-mini",

        temperature=0,

        response_format={
            "type": "json_object"
        },

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ]
    )

    result_text = (
        response
        .choices[0]
        .message
        .content
    )

    return json.loads(
        result_text
    )


# =========================================================
# DOCX GENERATOR
# =========================================================

def create_resume_docx(
    resume_data
):
    """
    Convert tailored resume JSON into
    ATS-friendly Word document.
    """

    document = Document()

    # -----------------------------------------------------
    # Summary
    # -----------------------------------------------------

    document.add_heading(
        "Professional Summary",
        level=1
    )

    document.add_paragraph(
        resume_data.get(
            "professional_summary",
            ""
        )
    )


    # -----------------------------------------------------
    # Skills
    # -----------------------------------------------------

    document.add_heading(
        "Core Skills",
        level=1
    )

    skills = resume_data.get(
        "skills",
        []
    )

    document.add_paragraph(
        " | ".join(skills)
    )


    # -----------------------------------------------------
    # Experience
    # -----------------------------------------------------

    document.add_heading(
        "Professional Experience",
        level=1
    )

    for item in resume_data.get(
        "experience",
        []
    ):

        role = item.get(
            "role",
            ""
        )

        company = item.get(
            "company",
            ""
        )

        heading = (
            f"{role} — {company}"
        )

        document.add_heading(
            heading,
            level=2
        )

        dates = item.get(
            "dates",
            ""
        )

        if dates:

            document.add_paragraph(
                dates
            )

        for bullet in item.get(
            "bullets",
            []
        ):

            document.add_paragraph(
                bullet,
                style="List Bullet"
            )


    # -----------------------------------------------------
    # Projects
    # -----------------------------------------------------

    projects = resume_data.get(
        "projects",
        []
    )

    if projects:

        document.add_heading(
            "Projects",
            level=1
        )

        for project in projects:

            document.add_heading(
                project.get(
                    "name",
                    ""
                ),
                level=2
            )

            for bullet in project.get(
                "bullets",
                []
            ):

                document.add_paragraph(
                    bullet,
                    style="List Bullet"
                )


    # -----------------------------------------------------
    # Education
    # -----------------------------------------------------

    education = resume_data.get(
        "education",
        []
    )

    if education:

        document.add_heading(
            "Education",
            level=1
        )

        for item in education:

            document.add_paragraph(
                str(item)
            )


    # -----------------------------------------------------
    # Certifications
    # -----------------------------------------------------

    certifications = resume_data.get(
        "certifications",
        []
    )

    if certifications:

        document.add_heading(
            "Certifications",
            level=1
        )

        for item in certifications:

            document.add_paragraph(
                str(item),
                style="List Bullet"
            )


    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


# =========================================================
# STREAMLIT CONFIGURATION
# =========================================================

st.set_page_config(

    page_title=
        "Agentic Career Copilot",

    page_icon="💼",

    layout="wide"
)


st.title(
    "💼 Agentic Career Copilot"
)

st.caption(
    "Discover → Evaluate → Tailor → "
    "Apply → Track → Prepare"
)

st.divider()


# =========================================================
# CANDIDATE KNOWLEDGE BASE
# =========================================================

st.subheader(
    "Candidate Knowledge Base"
)


if resume_exists():

    resume_version = (
        get_resume_hash()
    )

    st.success(
        "✓ Master resume is active"
    )

    st.caption(
        f"Resume version ID: "
        f"{resume_version}"
    )


    with open(
        MASTER_RESUME_PATH,
        "rb"
    ) as f:

        resume_text = (
            extract_pdf_text(f)
        )


    st.write(
        f"Resume contains approximately "
        f"**{len(resume_text.split())} words**."
    )


    with st.expander(
        "Preview extracted resume"
    ):

        st.text(
            resume_text[:4000]
        )


    replace_resume = (
        st.file_uploader(

            "Replace master resume",

            type=["pdf"],

            key="replace_resume"
        )
    )


    if replace_resume:

        save_master_resume(
            replace_resume
        )

        st.success(
            "Master resume replaced successfully."
        )

        st.rerun()


else:

    uploaded_resume = (
        st.file_uploader(

            "Upload your master resume",

            type=["pdf"]
        )
    )


    if uploaded_resume:

        save_master_resume(
            uploaded_resume
        )

        st.success(
            "Master resume saved."
        )

        st.rerun()


st.divider()


# =========================================================
# JOB DESCRIPTION
# =========================================================

st.subheader(
    "Job Description"
)


job_description = st.text_area(

    "Paste the job description",

    height=300,

    placeholder=
        "Paste the complete job description here..."
)


# =========================================================
# ANALYZE JOB
# =========================================================

if st.button(
    "Analyze Job Fit",
    type="primary"
):

    if not resume_exists():

        st.warning(
            "Upload your master resume first."
        )

        st.stop()


    if not job_description.strip():

        st.warning(
            "Paste a job description first."
        )

        st.stop()


    # -----------------------------------------------------
    # Load resume
    # -----------------------------------------------------

    with open(
        MASTER_RESUME_PATH,
        "rb"
    ) as f:

        resume_text = (
            extract_pdf_text(f)
        )


    resume_version = (
        get_resume_hash()
    )


    # -----------------------------------------------------
    # Index resume
    # -----------------------------------------------------

    with st.spinner(
        "Loading Candidate Knowledge Base..."
    ):

        chunk_count = (
            index_resume(
                resume_text,
                resume_version
            )
        )


    # -----------------------------------------------------
    # RAG retrieval
    # -----------------------------------------------------

    with st.spinner(
        "Retrieving relevant candidate evidence..."
    ):

        evidence = (
            retrieve_resume_evidence(
                job_description,
                top_k=6
            )
        )


    if not evidence:

        st.error(
            "No resume evidence could be retrieved."
        )

        st.stop()


    # -----------------------------------------------------
    # Fit Agent
    # -----------------------------------------------------

    with st.spinner(
        "Fit Agent is analyzing the role..."
    ):

        try:

            fit_result = (
                run_fit_agent(
                    job_description,
                    evidence
                )
            )

        except Exception as error:

            st.error(
                f"Fit Agent error: {error}"
            )

            st.stop()


    # =====================================================
    # FIT RESULTS
    # =====================================================

    st.divider()

    st.header(
        "Job Fit Analysis"
    )


    score = fit_result.get(
        "fit_score",
        0
    )


    col1, col2 = st.columns(
        [1, 3]
    )


    with col1:

        st.metric(
            "Overall Fit",
            f"{score}%"
        )


    with col2:

        st.progress(
            max(
                0,
                min(
                    score / 100,
                    1
                )
            )
        )

        st.write(
            fit_result.get(
                "summary",
                ""
            )
        )


    # =====================================================
    # STRONG MATCHES
    # =====================================================

    st.subheader(
        "✅ Strong Matches"
    )


    strong_matches = (
        fit_result.get(
            "strong_matches",
            []
        )
    )


    if strong_matches:

        for match in strong_matches:

            st.markdown(
                f"**{match.get('requirement', '')}**"
            )

            st.write(
                match.get(
                    "reason",
                    ""
                )
            )

            with st.expander(
                "View supporting evidence"
            ):

                st.write(
                    match.get(
                        "evidence",
                        ""
                    )
                )

    else:

        st.write(
            "No strong matches identified."
        )


    # =====================================================
    # PARTIAL MATCHES
    # =====================================================

    st.subheader(
        "🟡 Partial Matches"
    )


    partial_matches = (
        fit_result.get(
            "partial_matches",
            []
        )
    )


    if partial_matches:

        for match in partial_matches:

            st.markdown(
                f"**{match.get('requirement', '')}**"
            )

            st.write(
                match.get(
                    "reason",
                    ""
                )
            )

            with st.expander(
                "View supporting evidence"
            ):

                st.write(
                    match.get(
                        "evidence",
                        ""
                    )
                )

    else:

        st.write(
            "No partial matches identified."
        )


    # =====================================================
    # GAPS
    # =====================================================

    st.subheader(
        "🔴 Gaps"
    )


    gaps = fit_result.get(
        "gaps",
        []
    )


    if gaps:

        for gap in gaps:

            st.markdown(
                f"**{gap.get('requirement', '')}**"
            )

            st.write(
                gap.get(
                    "reason",
                    ""
                )
            )

    else:

        st.write(
            "No major gaps identified."
        )


    # =====================================================
    # FIT AGENT SUMMARY
    # =====================================================

    st.subheader(
        "Fit Agent Summary"
    )

    st.info(
        fit_result.get(
            "recommendation",
            ""
        )
    )


    # =====================================================
    # RAG TRANSPARENCY
    # =====================================================

    with st.expander(
        "🔎 RAG Evidence Retrieved"
    ):

        st.caption(
            f"Resume version: "
            f"{resume_version}"
        )

        st.caption(
            f"{chunk_count} resume chunks indexed"
        )


        for number, chunk in enumerate(
            evidence,
            start=1
        ):

            st.markdown(
                f"**Evidence {number}**"
            )

            st.write(
                chunk
            )

            st.divider()


    # =====================================================
    # SAVE ANALYSIS IN SESSION
    # =====================================================

    st.session_state[
        "fit_result"
    ] = fit_result

    st.session_state[
        "evidence"
    ] = evidence

    st.session_state[
        "job_description"
    ] = job_description

    st.session_state[
        "resume_text"
    ] = resume_text


# =========================================================
# RESUME AGENT
# =========================================================

if (
    "fit_result" in st.session_state
    and
    "evidence" in st.session_state
):

    st.divider()

    st.header(
        "Tailored Resume"
    )

    st.caption(
        "The Resume Agent can rephrase and prioritize "
        "existing experience but cannot invent "
        "unsupported experience."
    )


    if st.button(
        "Generate Tailored Resume"
    ):

        with st.spinner(
            "Resume Agent is tailoring your resume..."
        ):

            try:

                tailored_resume = (
                    run_resume_agent(

                        st.session_state[
                            "job_description"
                        ],

                        st.session_state[
                            "resume_text"
                        ],

                        st.session_state[
                            "evidence"
                        ],

                        st.session_state[
                            "fit_result"
                        ]
                    )
                )

                st.session_state[
                    "tailored_resume"
                ] = tailored_resume

            except Exception as error:

                st.error(
                    f"Resume Agent error: {error}"
                )

                st.stop()


# =========================================================
# DISPLAY TAILORED RESUME
# =========================================================

if "tailored_resume" in st.session_state:

    tailored_resume = (
        st.session_state[
            "tailored_resume"
        ]
    )


    st.success(
        "Tailored resume generated."
    )


    # -----------------------------------------------------
    # Summary
    # -----------------------------------------------------

    st.subheader(
        "Professional Summary"
    )

    st.write(
        tailored_resume.get(
            "professional_summary",
            ""
        )
    )


    # -----------------------------------------------------
    # Skills
    # -----------------------------------------------------

    st.subheader(
        "Skills"
    )

    skills = tailored_resume.get(
        "skills",
        []
    )

    st.write(
        " • ".join(
            skills
        )
    )


    # -----------------------------------------------------
    # Experience preview
    # -----------------------------------------------------

    st.subheader(
        "Professional Experience"
    )


    for experience in tailored_resume.get(
        "experience",
        []
    ):

        role = experience.get(
            "role",
            ""
        )

        company = experience.get(
            "company",
            ""
        )

        st.markdown(
            f"### {role} — {company}"
        )


        if experience.get(
            "dates"
        ):

            st.caption(
                experience.get(
                    "dates"
                )
            )


        for bullet in experience.get(
            "bullets",
            []
        ):

            st.write(
                f"• {bullet}"
            )


    # -----------------------------------------------------
    # Projects preview
    # -----------------------------------------------------

    projects = tailored_resume.get(
        "projects",
        []
    )


    if projects:

        st.subheader(
            "Projects"
        )


        for project in projects:

            st.markdown(
                f"### "
                f"{project.get('name', '')}"
            )


            for bullet in project.get(
                "bullets",
                []
            ):

                st.write(
                    f"• {bullet}"
                )


    # -----------------------------------------------------
    # Unsupported requirements
    # -----------------------------------------------------

    unsupported = (
        tailored_resume.get(
            "unsupported_requirements",
            []
        )
    )


    if unsupported:

        with st.expander(
            "⚠️ JD requirements not added to resume"
        ):

            st.caption(
                "These requirements were not supported "
                "by evidence in the master resume."
            )


            for item in unsupported:

                st.write(
                    f"• {item}"
                )


    # -----------------------------------------------------
    # DOCX
    # -----------------------------------------------------

    docx_file = (
        create_resume_docx(
            tailored_resume
        )
    )


    st.download_button(

        label=
            "⬇️ Download Tailored Resume (.docx)",

        data=
            docx_file.getvalue(),

        file_name=
            "tailored_resume.docx",

        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        )
    )